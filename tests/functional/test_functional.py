import pytest

import os
import pytest
from flask import url_for

@pytest.fixture
def test_files():
    # Create a sample Word file and Excel file for testing
    word_file_path = os.path.join(os.path.dirname(__file__), '..', 'static', 'sample.docx')
    excel_file_path = os.path.join(os.path.dirname(__file__), '..', 'static', 'sample.xlsx')
    
    # Create a simple Word file with placeholders
    from docx import Document
    doc = Document()
    doc.add_paragraph('Hello {{name}}!')
    doc.save(word_file_path)
    
    # Create a simple Excel file with data
    import pandas as pd
    df = pd.DataFrame({'name': ['Alice', 'Bob']})
    df.to_excel(excel_file_path, index=False)
    
    yield word_file_path, excel_file_path

def test_dashboard_manual_upload(client, test_files):
    word_file_path, _ = test_files
    
    with open(word_file_path, 'rb') as word_file:
        response = client.post('/dashboard', data={
            'wordFile': word_file,
            'radioButton': 'manually',
            'batchQuantity': '2'
        }, content_type='multipart/form-data')
    
    assert response.status_code == 200
    assert b"Uploaded file:" in response.data
    assert b"PDF 1:" in response.data

def test_dashboard_excel_upload(client, test_files):
    word_file_path, excel_file_path = test_files
    
    with open(word_file_path, 'rb') as word_file, open(excel_file_path, 'rb') as excel_file:
        response = client.post('/dashboard', data={
            'wordFile': word_file,
            'radioButton': 'excel',
            'excelFile': excel_file
        }, content_type='multipart/form-data')
    
    assert response.status_code == 200
    assert b"Uploaded file:" in response.data
    assert b"PDF 1:" in response.data

def test_generate_pdf_preview(client, test_files):
    word_file_path, _ = test_files
    
    with open(word_file_path, 'rb') as word_file:
        client.post('/dashboard', data={
            'wordFile': word_file,
            'radioButton': 'manually',
            'batchQuantity': '1'
        }, content_type='multipart/form-data')
    
    response = client.post('/generate-pdf-preview', json={
        "index": 0,
        "formData": {
            "0_name": "Alice"
        }
    })
    
    assert response.status_code == 200
    assert response.json['message'] == "PDF generated successfully"

# def test_transform_file(client, test_files):
#     word_file_path, _ = test_files
    
#     with open(word_file_path, 'rb') as word_file:
#         client.post('/dashboard', data={
#             'wordFile': word_file,
#             'radioButton': 'manually',
#             'batchQuantity': '1'
#         }, content_type='multipart/form-data')
    
#     response = client.post('/post-doc-form', json={
#         "formDataArray": [{"name": "Alice"}],
#         "fileName": "name"
#     })
    
#     assert response.status_code == 200
#     assert response.headers['Content-Type'] == 'application/zip'

# def test_invalid_word_file_upload(client):
#     response = client.post('/dashboard', data={
#         'wordFile': (b'fake content', 'test.txt'),
#         'radioButton': 'manually',
#         'batchQuantity': '1'
#     }, content_type='multipart/form-data')
    
#     assert response.status_code == 400
#     assert b"Invalid file format" in response.data

# def test_invalid_excel_file_upload(client, test_files):
#     word_file_path, _ = test_files
    
#     with open(word_file_path, 'rb') as word_file:
#         response = client.post('/dashboard', data={
#             'wordFile': word_file,
#             'radioButton': 'excel',
#             'excelFile': (b'fake content', 'test.txt')
#         }, content_type='multipart/form-data')
    
#     assert response.status_code == 400
#     assert b"Invalid Excel file format" in response.data


# def test_save_file(client):
#     # Simulate filling the form and saving the file
#     with open('tests/sample.docx', 'rb') as f:
#         client.post('/dashboard', data={'wordFile': f, 'radioButton': 'manually', 'batchQuantity': '1'})
#     response = client.post('/post-doc-form', json={
#         'formDataArray': [{'placeholder1': 'value1', 'placeholder2': 'value2'}],
#         'fileName': 'test_file'
#     })
#     assert response.status_code == 200
#     assert response.headers['Content-Type'] == 'application/zip'  # Check if the response is a ZIP file

# def test_word_file_upload(client, sample_docx_file):
#     # Test valid .docx file upload
#     with open(sample_docx_file, 'rb') as f:
#         response = client.post('/dashboard', data={'wordFile': f, 'radioButton': 'manually', 'batchQuantity': '1'})
#     assert response.status_code == 200
#     assert b"File example.docx uploaded successfully!" in response.data

#     # Test invalid file upload
#     with open('../static/wrong.txt', 'rb') as f:
#         response = client.post('/dashboard', data={'wordFile': f, 'radioButton': 'manually', 'batchQuantity': '1'})
#     assert response.status_code == 400
#     assert b"Invalid file type. Please upload a valid Word (.doc or .docx) file." in response.data